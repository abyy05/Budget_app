
# Finished code
# slight problem in import: delete the id column and stat from the name column. then import data as usual
 
import tkinter as tk
from tkinter import messagebox, ttk, filedialog
import sqlite3
import pandas as pd
import openpyxl
import xlsxwriter
# Database Setup

conn = sqlite3.connect("budget.db")
cur = conn.cursor()

cur.executescript('''
BEGIN;

CREATE TABLE IF NOT EXISTS income (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    amount REAL NOT NULL
);

CREATE TABLE IF NOT EXISTS expense (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    category TEXT NOT NULL,
    amount REAL NOT NULL
);

CREATE TABLE IF NOT EXISTS saving (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    name TEXT NOT NULL,
    amount REAL NOT NULL
);

COMMIT;
''')
conn.commit()

# Helper Functions
def is_valid_name(value):
    return not value.strip().isdigit() and value.strip() != ""

def is_valid_amount(value):
    try:
        float(value)
        return True
    except ValueError:
        return False

# Helper Functions
def get_total(table):
    cur.execute(f"SELECT SUM(amount) FROM {table}")
    result = cur.fetchone()[0]
    return result if result else 0

def update_budget_status():
    income_total = get_total('income')
    expense_total = get_total('expense')
    saving_total = get_total('saving')
    budget = income_total - (expense_total + saving_total)

    if income_total == 0 and expense_total == 0 and saving_total == 0:
        budget_status_label.config(text="📓 Empty notebook. No data entered yet.", fg="gray")   
    elif budget == 0:
        budget_status_label.config(text="✅ Budget Completed: Well done!", fg="green")
    elif budget<0:
        budget_status_label.config(text=f"⚠️ Budget limit crossed: ₹{budget:.2f} in shortage. Redo budget", fg="red")
    else:
        budget_status_label.config(text=f"⚠️ Budget incomplete: ₹{budget:.2f} left to allocate.", fg="purple")

def show_total(table_name):
    total = get_total(table_name)
    labels = {
        "income": "💰 Total Income",
        "expense": "💸 Total Expenses",
        "saving": "🏦 Total Savings"
    }
    messagebox.showinfo(labels[table_name], f"{labels[table_name]}: ₹{total:.2f}")

def refresh_tables():
    for tree, table in [(income_tree, 'income'), (expense_tree, 'expense'), (saving_tree, 'saving')]:
        for i in tree.get_children():
            tree.delete(i)
        cur.execute(f"SELECT * FROM {table}")
        for row in cur.fetchall():
            tree.insert('', 'end', values=row)
    update_budget_status()

def add_entry(table, entries):
    data = tuple(e.get() for e in entries)

    if table == 'expense':
        name, category, amount = data
        if not is_valid_name(name) or not is_valid_name(category):
            messagebox.showerror("Invalid Input", "Name and Category should not be just numbers or empty.")
            return
    else:
        name, amount = data
        if not is_valid_name(name):
            messagebox.showerror("Invalid Input", "Name should not be just numbers or empty.")
            return

    if not is_valid_amount(amount):
        messagebox.showerror("Invalid Input", "Amount must be a valid number.")
        return

    placeholders = ', '.join('?' * len(data))
    with conn:
        cur.execute(f"INSERT INTO {table} (name, {('category, ' if table == 'expense' else '')}amount) VALUES ({placeholders})", data)
    for e in entries:
        e.delete(0, tk.END)
    refresh_tables()

def delete_entry(table, tree):
    selected = tree.selection()
    if not selected:
        messagebox.showerror("Error", "Select an entry to delete.")
        return
    item = tree.item(selected[0])['values']
    with conn:
        cur.execute(f"DELETE FROM {table} WHERE id=?", (item[0],))
    refresh_tables()

def clear_table(table):
    confirm = messagebox.askyesno("Confirm Clear", f"Are you sure you want to delete all data from {table}?")
    if confirm:
        with conn:
            cur.execute(f"DELETE FROM {table}")
        refresh_tables()

def export_table(table_name):
    cur.execute(f"SELECT * FROM {table_name}")
    rows = cur.fetchall()
    columns = [desc[0] for desc in cur.description]

    df = pd.DataFrame(rows, columns=columns)

    file_path = filedialog.asksaveasfilename(
        defaultextension=".csv",
        filetypes=[("CSV files", "*.csv"), ("Excel files", "*.xlsx")],
        title="Save As"
    )
    if file_path:
        try:
            if file_path.endswith('.csv'):
                df.to_csv(file_path, index=False)
            else:
                df.to_excel(file_path, index=False)
            messagebox.showinfo("Export Success", f"{table_name.capitalize()} exported successfully!")
        except Exception as e:
            messagebox.showerror("Export Error", str(e))

def import_table(table_name):
    expected_cols = {
        "income": {"id","name", "amount"},
        "expense": {"id","name", "category", "amount"},
        "saving": {"id","name", "amount"}
    }

    file_path = filedialog.askopenfilename(
        filetypes=[("CSV/Excel files", "*.csv *.xlsx")],
        title="Select File to Import"
    )
    if not file_path:
        return

    try:
        if file_path.endswith('.csv'):
            df = pd.read_csv(file_path)
        else:
            df = pd.read_excel(file_path)

        if set(df.columns) != expected_cols[table_name]:
            raise ValueError(f"Incorrect columns. Expected: {expected_cols[table_name]}")

        cols = ', '.join(df.columns)
        with conn:
            for _, row in df.iterrows():
                values = tuple(row[col] for col in df.columns)
                placeholders = ', '.join('?' * len(values))
                cur.execute(f"INSERT INTO {table_name} ({cols}) VALUES ({placeholders})", values)

        refresh_tables()
        messagebox.showinfo("Import Success", f"{table_name.capitalize()} data imported successfully!")

    except Exception as e:
        messagebox.showerror("Import Failed", str(e))

# GUI Setup
root = tk.Tk()
root.title("Zero-Based Budgeting App")

# --- Income Frame ---
income_frame = tk.LabelFrame(root, text="Income")
income_frame.grid(row=0, column=0, padx=10, pady=5)
tk.Label(income_frame, text="Name").grid(row=0, column=0)
tk.Label(income_frame, text="Amount").grid(row=0, column=1)
income_name = tk.Entry(income_frame)
income_amount = tk.Entry(income_frame)
income_name.grid(row=1, column=0)
income_amount.grid(row=1, column=1)
tk.Button(income_frame, text="Add Income", command=lambda: add_entry('income', [income_name, income_amount])).grid(row=1, column=2)
income_tree = ttk.Treeview(income_frame, columns=("ID", "Name", "Amount"), show='headings')
for col in income_tree['columns']:
    income_tree.heading(col, text=col)
income_tree.grid(row=2, column=0, columnspan=3)
tk.Button(income_frame, text="Delete Selected", command=lambda: delete_entry('income', income_tree)).grid(row=3, column=0)
tk.Button(income_frame, text="Show Total", command=lambda: show_total('income')).grid(row=3, column=1)
tk.Button(income_frame, text="Clear All", command=lambda: clear_table('income')).grid(row=3, column=2)
tk.Button(income_frame, text="Export", command=lambda: export_table('income')).grid(row=4, column=0)
tk.Button(income_frame, text="Import", command=lambda: import_table('income')).grid(row=4, column=1)

# --- Expense Frame ---
expense_frame = tk.LabelFrame(root, text="Expenses")
expense_frame.grid(row=0, column=1, padx=10, pady=5)
tk.Label(expense_frame, text="Name").grid(row=0, column=0)
tk.Label(expense_frame, text="Category").grid(row=0, column=1)
tk.Label(expense_frame, text="Amount").grid(row=0, column=2)
exp_name = tk.Entry(expense_frame)
exp_category = tk.Entry(expense_frame)
exp_amount = tk.Entry(expense_frame)
exp_name.grid(row=1, column=0)
exp_category.grid(row=1, column=1)
exp_amount.grid(row=1, column=2)
tk.Button(expense_frame, text="Add Expense", command=lambda: add_entry('expense', [exp_name, exp_category, exp_amount])).grid(row=1, column=3)
expense_tree = ttk.Treeview(expense_frame, columns=("ID", "Name", "Category", "Amount"), show='headings')
for col in expense_tree['columns']:
    expense_tree.heading(col, text=col)
expense_tree.grid(row=2, column=0, columnspan=4)
tk.Button(expense_frame, text="Delete Selected", command=lambda: delete_entry('expense', expense_tree)).grid(row=3, column=0)
tk.Button(expense_frame, text="Show Total", command=lambda: show_total('expense')).grid(row=3, column=2)
tk.Button(expense_frame, text="Clear All", command=lambda: clear_table('expense')).grid(row=3, column=3)
tk.Button(expense_frame, text="Export", command=lambda: export_table('expense')).grid(row=4, column=0)
tk.Button(expense_frame, text="Import", command=lambda: import_table('expense')).grid(row=4, column=3)

# --- Saving Frame ---
saving_frame = tk.LabelFrame(root, text="Savings")
saving_frame.grid(row=1, column=0, padx=10, pady=5)
tk.Label(saving_frame, text="Name").grid(row=0, column=0)
tk.Label(saving_frame, text="Amount").grid(row=0, column=1)
save_name = tk.Entry(saving_frame)
save_amount = tk.Entry(saving_frame)
save_name.grid(row=1, column=0)
save_amount.grid(row=1, column=1)
tk.Button(saving_frame, text="Add Saving", command=lambda: add_entry('saving', [save_name, save_amount])).grid(row=1, column=2)
saving_tree = ttk.Treeview(saving_frame, columns=("ID", "Name", "Amount"), show='headings')
for col in saving_tree['columns']:
    saving_tree.heading(col, text=col)
saving_tree.grid(row=2, column=0, columnspan=3)
tk.Button(saving_frame, text="Delete Selected", command=lambda: delete_entry('saving', saving_tree)).grid(row=3, column=0)
tk.Button(saving_frame, text="Show Total", command=lambda: show_total('saving')).grid(row=3, column=1)
tk.Button(saving_frame, text="Clear All", command=lambda: clear_table('saving')).grid(row=3, column=2)
tk.Button(saving_frame, text="Export", command=lambda: export_table('saving')).grid(row=4, column=0)
tk.Button(saving_frame, text="Import", command=lambda: import_table('saving')).grid(row=4, column=2)

# --- Formula and Budget Status ---
formula_label = tk.Label(root, text="Formula: Income - (Expenses + Savings) = Zero Budget", font=("Arial", 11, "italic"))
formula_label.grid(row=2, column=0, columnspan=2)

budget_status_label = tk.Label(root, text="📊 Budget status will appear here.", font=("Arial", 12, "bold"))
budget_status_label.grid(row=3, column=0, columnspan=2, pady=10)

refresh_tables()
root.mainloop()
